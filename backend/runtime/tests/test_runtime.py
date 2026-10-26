"""Real HTTP/database tests against an isolated temporary workspace."""

import io
import os
import tempfile
import time
import uuid

import pytest
from fastapi.testclient import TestClient

# Set before importing the app: never run these tests against user data.
TEST_DATA = tempfile.TemporaryDirectory(prefix="rag-runtime-tests-")
os.environ["RAG_DATA_DIR"] = TEST_DATA.name
os.environ["RAG_PROVIDER"] = "extractive"

from backend.runtime.main import app, attempts  # noqa: E402

ORIGIN = {"origin": "http://localhost:3000"}
PASSWORD = "Long-enough-test-password!"


@pytest.fixture
def client():
    attempts.clear()
    with TestClient(app, headers=ORIGIN) as instance:
        yield instance


def register(client, name="Tester"):
    email = f"{uuid.uuid4().hex}@example.com"
    response = client.post(
        "/api/v2/auth/register",
        json={
            "email": email,
            "password": PASSWORD,
            "name": name,
            "organization": name + " organization",
        },
    )
    assert response.status_code == 201, response.text
    return response.json(), email, client.cookies.get("rag_session")


def upload_index(
    client,
    name="handbook.txt",
    content=b"Employees receive twenty days of annual leave. Submit leave requests to the HR manager.",
):
    response = client.post("/api/v2/documents", files={"file": (name, content)})
    assert response.status_code == 202, response.text
    doc_id = response.json()["id"]
    for _ in range(60):
        doc = client.get(f"/api/v2/documents/{doc_id}").json()
        if doc["status"] in {"indexed", "failed"}:
            break
        time.sleep(0.1)
    assert doc["status"] == "indexed", doc
    return doc_id


def conversation(client):
    result = client.post("/api/v2/conversations", json={})
    assert result.status_code == 201
    return result.json()["id"]


def test_authentication_and_revocation(client):
    _, email, token = register(client)
    assert (
        "HttpOnly"
        in client.post("/api/v2/auth/login", json={"email": email, "password": PASSWORD}).headers[
            "set-cookie"
        ]
    )
    assert (
        client.post(
            "/api/v2/auth/login", json={"email": email, "password": "incorrect-password"}
        ).status_code
        == 401
    )
    assert client.get("/api/v2/me").status_code == 200
    revoked = client.cookies.get("rag_session")
    assert client.post("/api/v2/auth/logout").status_code == 200
    client.cookies.set("rag_session", revoked)
    assert client.get("/api/v2/me").status_code == 401
    client.cookies.set("rag_session", "1")
    assert client.get("/api/v2/me").status_code == 401
    # Login rotates the previous session as well.
    client.cookies.set("rag_session", token)
    assert client.get("/api/v2/me").status_code == 401


def test_tenant_isolation_and_forged_context(client):
    _, _, owner_token = register(client, "Organization A")
    document_id = upload_index(client)
    conv_id = conversation(client)
    register(client, "Organization B")
    client.headers["X-Tenant-ID"] = "Organization A"
    client.headers["X-Organization-ID"] = "Organization A"
    assert client.get(f"/api/v2/documents/{document_id}").status_code == 404
    assert client.get(f"/api/v2/documents/{document_id}/download").status_code == 404
    assert client.delete(f"/api/v2/documents/{document_id}").status_code == 404
    assert client.get(f"/api/v2/conversations/{conv_id}/messages").status_code == 404
    assert (
        client.post(
            f"/api/v2/conversations/{conv_id}/messages", json={"query": "annual leave"}
        ).status_code
        == 404
    )
    assert (
        client.post(
            "/api/v2/search", json={"query": "annual leave", "doc_ids": [document_id]}
        ).json()
        == []
    )
    assert client.get("/api/v2/documents").json() == []
    client.cookies.set("rag_session", owner_token)
    assert client.get(f"/api/v2/documents/{document_id}").status_code == 200


def test_upload_validation_duplicate_preview_and_deletion(client):
    register(client)
    assert (
        client.post("/api/v2/documents", files={"file": ("fake.pdf", b"not a PDF")}).status_code
        == 422
    )
    assert (
        client.post("/api/v2/documents", files={"file": ("fake.docx", b"not a zip")}).status_code
        == 422
    )
    assert (
        client.post("/api/v2/documents", files={"file": ("script.exe", b"bad")}).status_code == 422
    )
    assert (
        client.post("/api/v2/documents", files={"file": ("binary.txt", b"a\0b")}).status_code == 422
    )
    doc_id = upload_index(client)
    duplicate = client.post(
        "/api/v2/documents",
        files={
            "file": (
                "same.txt",
                b"Employees receive twenty days of annual leave. Submit leave requests to the HR manager.",
            )
        },
    )
    assert duplicate.status_code == 409
    doc = client.get(f"/api/v2/documents/{doc_id}").json()
    assert "twenty days" in doc["sections"][0]["content"]
    assert doc["sections"][0]["page"] is None
    assert client.get(f"/api/v2/documents/{doc_id}/download").status_code == 200
    assert client.delete(f"/api/v2/documents/{doc_id}").status_code == 200
    assert client.post("/api/v2/search", json={"query": "annual leave"}).json() == []
    assert client.get(f"/api/v2/documents/{doc_id}/download").status_code == 404


def test_stream_persistence_citations_and_analytics(client):
    register(client)
    doc_id = upload_index(client)
    conv_id = conversation(client)
    response = client.post(
        f"/api/v2/conversations/{conv_id}/messages", json={"query": "annual leave"}
    )
    assert response.headers["content-type"].startswith("text/event-stream")
    assert all(f"event: {name}" in response.text for name in ["meta", "sources", "token", "done"])
    assert "no AI model configured" in response.text
    assert "twenty days" in response.text
    saved = client.get(f"/api/v2/conversations/{conv_id}/messages").json()
    assert len(saved) == 2 and saved[1]["status"] == "completed"
    assert saved[1]["sources"][0]["doc_id"] == doc_id
    assert client.get("/api/v2/analytics").json()["queries"] == 1
    client.delete(f"/api/v2/documents/{doc_id}")
    assert client.get(f"/api/v2/conversations/{conv_id}/messages").json()[1]["sources"] == []


def test_provider_stream_success_error_and_invalid_citation(client, monkeypatch):
    from backend.runtime import providers

    register(client)
    upload_index(client)
    monkeypatch.setattr(
        providers,
        "configuration",
        lambda: {"provider": "test", "configured": True, "mode": "generation", "model": "test"},
    )

    async def stream(*args):
        yield "Twenty days "
        yield "of leave [1]."

    monkeypatch.setattr(providers, "generate", stream)
    conv_id = conversation(client)
    result = client.post(
        f"/api/v2/conversations/{conv_id}/messages", json={"query": "annual leave"}
    )
    assert result.text.count("event: token") == 2
    assert (
        client.get(f"/api/v2/conversations/{conv_id}/messages").json()[1]["content"]
        == "Twenty days of leave [1]."
    )

    async def invalid(*args):
        yield "Invented citation [99]"

    monkeypatch.setattr(providers, "generate", invalid)
    conv_id = conversation(client)
    result = client.post(
        f"/api/v2/conversations/{conv_id}/messages", json={"query": "annual leave"}
    )
    assert "event: error" in result.text
    assert client.get(f"/api/v2/conversations/{conv_id}/messages").json()[1]["status"] == "failed"

    async def broken(*args):
        raise TimeoutError("provider failed")
        yield "unreachable"

    monkeypatch.setattr(providers, "generate", broken)
    conv_id = conversation(client)
    result = client.post(
        f"/api/v2/conversations/{conv_id}/messages", json={"query": "annual leave"}
    )
    assert "event: error" in result.text and "event: done" in result.text


def test_member_permissions_and_immediate_deactivation(client):
    _, admin_email, _ = register(client)
    email = f"viewer-{uuid.uuid4().hex}@example.com"
    member = client.post(
        "/api/v2/admin/users",
        json={"name": "Viewer", "email": email, "password": PASSWORD, "role": "viewer"},
    ).json()
    assert (
        client.post("/api/v2/auth/login", json={"email": email, "password": PASSWORD}).status_code
        == 200
    )
    viewer_token = client.cookies.get("rag_session")
    assert client.get("/api/v2/admin/users").status_code == 403
    assert (
        client.post("/api/v2/documents", files={"file": ("test.txt", b"some text")}).status_code
        == 403
    )
    # Logging into the viewer account rotates this browser's previous admin session.
    assert (
        client.post(
            "/api/v2/auth/login", json={"email": admin_email, "password": PASSWORD}
        ).status_code
        == 200
    )
    assert (
        client.patch(
            f"/api/v2/admin/users/{member['id']}", json={"active": False, "role": "viewer"}
        ).status_code
        == 200
    )
    client.cookies.set("rag_session", viewer_token)
    assert client.get("/api/v2/me").status_code == 401
    assert (
        client.post("/api/v2/auth/login", json={"email": email, "password": PASSWORD}).status_code
        == 401
    )


def test_origin_and_limits(client):
    assert (
        client.post(
            "/api/v2/auth/login",
            headers={"origin": "https://evil.example"},
            json={"email": "x@example.com", "password": PASSWORD},
        ).status_code
        == 403
    )
    for _ in range(15):
        client.post(
            "/api/v2/auth/login", json={"email": "missing@example.com", "password": PASSWORD}
        )
    assert (
        client.post(
            "/api/v2/auth/login", json={"email": "missing@example.com", "password": PASSWORD}
        ).status_code
        == 429
    )


def test_real_binary_parsers(client):
    import pymupdf
    from docx import Document
    from openpyxl import Workbook

    register(client)
    pdf = pymupdf.open()
    pdf.new_page().insert_text((72, 72), "PDF payroll policy: expense reports are due Friday.")
    doc_id = upload_index(client, "policy.pdf", pdf.tobytes())
    assert client.get(f"/api/v2/documents/{doc_id}").json()["sections"][0]["page"] == 1
    pdf.close()
    word = Document()
    word.add_paragraph("DOCX procedure: equipment requests require manager approval.")
    word_bytes = io.BytesIO()
    word.save(word_bytes)
    upload_index(client, "procedure.docx", word_bytes.getvalue())
    book = Workbook()
    book.active.append(["Department", "Budget"])
    book.active.append(["Engineering", 5000])
    sheet_bytes = io.BytesIO()
    book.save(sheet_bytes)
    book.close()
    upload_index(client, "budget.xlsx", sheet_bytes.getvalue())


def test_no_evidence_and_chat_rename_delete(client):
    register(client)
    conv_id = conversation(client)
    result = client.post(
        f"/api/v2/conversations/{conv_id}/messages", json={"query": "unanswerable question"}
    )
    assert "could not find supporting evidence" in result.text
    assert (
        client.patch(f"/api/v2/conversations/{conv_id}", json={"title": "Renamed"}).status_code
        == 200
    )
    assert client.get("/api/v2/conversations").json()[0]["title"] == "Renamed"
    assert client.delete(f"/api/v2/conversations/{conv_id}").status_code == 200
    assert client.get(f"/api/v2/conversations/{conv_id}/messages").status_code == 404


def test_chunk_boundaries():
    from backend.runtime.ingestion import split_text

    text = "This is a very long paragraph. " * 200
    chunks = list(split_text(text))
    assert len(chunks) > 1
    assert all(len(chunk) <= 1400 for chunk in chunks)
    assert chunks[-1].endswith("paragraph.")
    with pytest.raises(ValueError):
        list(split_text(text, size=100, overlap=100))
