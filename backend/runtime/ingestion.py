"""Bounded parsing and durable single-process ingestion; no synthetic progress."""

import asyncio
import io
import json
import zipfile
from html.parser import HTMLParser
from pathlib import Path

from sqlalchemy import delete, select

from .database import DATA, Chunk, Document, Session

SUPPORTED = {
    ".pdf",
    ".docx",
    ".xlsx",
    ".txt",
    ".md",
    ".csv",
    ".json",
    ".html",
    ".xml",
    ".py",
    ".ts",
    ".js",
}
MAX_BYTES = 20 * 1024 * 1024
MAX_TEXT = 2_000_000


class TextHTML(HTMLParser):
    def __init__(self):
        super().__init__()
        self.parts = []
        self.hidden = 0

    def handle_starttag(self, tag, attrs):
        if tag in {"script", "style"}:
            self.hidden += 1

    def handle_endtag(self, tag):
        if tag in {"script", "style"}:
            self.hidden = max(0, self.hidden - 1)

    def handle_data(self, data):
        if not self.hidden:
            self.parts.append(data)


def validate(data: bytes, suffix: str):
    if suffix not in SUPPORTED:
        raise ValueError("Unsupported format. Use PDF, DOCX, XLSX or UTF-8 text files.")
    if not data or len(data) > MAX_BYTES:
        raise ValueError("Files must contain data and be at most 20 MB.")
    if suffix == ".pdf" and not data.startswith(b"%PDF-"):
        raise ValueError("File content does not match PDF format.")
    if suffix in {".docx", ".xlsx"}:
        if not zipfile.is_zipfile(io.BytesIO(data)):
            raise ValueError("Invalid Office document.")
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            if sum(f.file_size for f in archive.infolist()) > 80 * 1024 * 1024:
                raise ValueError("Expanded Office document exceeds the parsing limit.")
            required = "word/document.xml" if suffix == ".docx" else "xl/workbook.xml"
            if required not in archive.namelist():
                raise ValueError("File content does not match its Office extension.")
    elif suffix != ".pdf":
        text = data.decode("utf-8-sig")
        if "\0" in text:
            raise ValueError("Binary content cannot be uploaded as text.")


def extract(path: Path, suffix: str) -> list[tuple[int | None, str]]:
    if suffix == ".pdf":
        import pymupdf

        with pymupdf.open(path) as pdf:
            if len(pdf) > 1000:
                raise ValueError("PDF exceeds the 1000-page parsing limit.")
            pages = [(i + 1, page.get_text()) for i, page in enumerate(pdf)]
    elif suffix == ".docx":
        from docx import Document as Word

        doc = Word(path)
        text = "\n".join(p.text for p in doc.paragraphs)
        text += "\n" + "\n".join(
            " | ".join(c.text for c in row.cells) for t in doc.tables for row in t.rows
        )
        pages = [(None, text)]
    elif suffix == ".xlsx":
        from openpyxl import load_workbook

        # Originals use UUID storage keys without extensions; openpyxl needs a file object.
        book = load_workbook(io.BytesIO(path.read_bytes()), read_only=True, data_only=True)
        try:
            pages = []
            for sheet in book:
                lines = [sheet.title]
                for index, row in enumerate(sheet.iter_rows(values_only=True)):
                    if index >= 20000:
                        raise ValueError("Spreadsheet exceeds 20,000 rows per sheet.")
                    lines.append(" | ".join(str(c) if c is not None else "" for c in row))
                pages.append((None, "\n".join(lines)))
        finally:
            book.close()
    else:
        text = path.read_text(encoding="utf-8-sig")
        if suffix == ".html":
            parser = TextHTML()
            parser.feed(text)
            text = "\n".join(parser.parts)
        if suffix == ".json":
            text = json.dumps(json.loads(text), ensure_ascii=False, indent=2)
        pages = [(None, text)]
    if sum(len(text) for _, text in pages) > MAX_TEXT:
        raise ValueError("Extracted text exceeds the two-million-character limit.")
    if not any(text.strip() for _, text in pages):
        raise ValueError("No text found. Scanned documents require OCR, which is not configured.")
    return pages


def split_text(text: str, size=1400, overlap=180):
    if size <= overlap or overlap < 0:
        raise ValueError("Chunk size must exceed overlap.")
    start = 0
    while start < len(text):
        end = min(start + size, len(text))
        if end < len(text):
            boundary = text.rfind(" ", start + size // 2, end)
            if boundary > start:
                end = boundary
        value = text[start:end].strip()
        if value:
            yield value
        if end == len(text):
            break
        start = end - overlap


async def process_document(doc_id: str):
    async with Session() as db:
        doc = await db.get(Document, doc_id)
        if not doc or doc.status != "queued":
            return
        doc.status = "extracting"
        await db.commit()
        path = DATA / "uploads" / doc.id
        try:
            pages = await asyncio.to_thread(extract, path, Path(doc.name).suffix.lower())
            # Reload after extraction: deletion can race with a parser running off-thread.
            await db.refresh(doc)
            doc.status = "chunking"
            await db.commit()
            await db.execute(
                delete(Chunk).where(Chunk.doc_id == doc.id, Chunk.org_id == doc.org_id)
            )
            position = 0
            for page, text in pages:
                for content in split_text(text):
                    db.add(
                        Chunk(
                            doc_id=doc.id,
                            org_id=doc.org_id,
                            position=position,
                            page=page,
                            content=content,
                        )
                    )
                    position += 1
            doc.status = "indexed"
            doc.error = None
            await db.commit()
        except Exception as exc:
            await db.rollback()
            doc = await db.get(Document, doc_id)
            if doc:
                doc.status = "failed"
                doc.error = (
                    str(exc)[:300]
                    if isinstance(exc, ValueError)
                    else "Document parsing failed. Check the file and retry."
                )
                await db.commit()


async def worker():
    # This runtime deliberately runs one worker; use a leased distributed queue before scaling.
    async with Session() as db:
        for doc in (
            await db.scalars(
                select(Document).where(Document.status.in_(["extracting", "chunking"]))
            )
        ).all():
            doc.status = "queued"
        await db.commit()
    while True:
        async with Session() as db:
            ids = (
                await db.scalars(select(Document.id).where(Document.status == "queued").limit(5))
            ).all()
        for doc_id in ids:
            await process_document(doc_id)
        await asyncio.sleep(1)
